#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Kalan tüm ünite sunumlarını hızlıca oluştur
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_basic_ppt(title, subtitle, slides_content, filename, color_rgb=(220, 100, 180)):
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # Kapak
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title_shape = slide.shapes.title
    subtitle_shape = slide.placeholders[1]
    title_shape.text = title
    subtitle_shape.text = subtitle
    title_shape.text_frame.paragraphs[0].font.size = Pt(54)
    title_shape.text_frame.paragraphs[0].font.bold = True
    title_shape.text_frame.paragraphs[0].font.color.rgb = RGBColor(*color_rgb)

    # İçerik slaytları
    for slide_title, content_list in slides_content:
        slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(slide_layout)
        title_shape = slide.shapes.title
        title_shape.text = slide_title
        title_shape.text_frame.paragraphs[0].font.size = Pt(36)
        title_shape.text_frame.paragraphs[0].font.color.rgb = RGBColor(*color_rgb)

        body_shape = slide.placeholders[1]
        tf = body_shape.text_frame
        tf.clear()
        for item in content_list:
            p = tf.add_paragraph()
            p.text = item
            p.font.size = Pt(20)

    prs.save(filename)
    print(f"✅ Oluşturuldu: {filename}")

# 5. Sınıf Ünite 2: Dijital Ürün Tasarımı
create_basic_ppt(
    "ÜNİTE 2: DİJİTAL ÜRÜN TASARIMI",
    "5. Sınıf - 6 Hafta",
    [
        ("Microsoft Word", ["Belge oluşturma", "Yazı biçimlendirme", "Resim ekleme", "Tablo oluşturma"]),
        ("Paint", ["Temel çizim araçları", "Renk seçimi", "Şekiller", "Metin ekleme"]),
        ("PowerPoint", ["Slayt oluşturma", "Tasarım şablonları", "Animasyonlar", "Sunum teknikleri"]),
    ],
    "/home/user/BST/sunumlar/5-sinif-unite2-dijital-urun.pptx",
    (100, 180, 220)
)

# 5. Sınıf Ünite 3: Ağlar ve İletişim
create_basic_ppt(
    "ÜNİTE 3: AĞLAR VE İLETİŞİM",
    "5. Sınıf - 3 Hafta",
    [
        ("İnternet", ["İnternet nedir?", "Tarayıcılar", "Arama motorları", "Güvenli internet"]),
        ("E-posta", ["E-posta kullanımı", "E-posta etiği", "Ek dosya gönderme"]),
    ],
    "/home/user/BST/sunumlar/5-sinif-unite3-aglar-iletisim.pptx",
    (180, 220, 100)
)

# 5. Sınıf Ünite 4: Etik ve Güvenlik
create_basic_ppt(
    "ÜNİTE 4: ETİK VE GÜVENLİK",
    "5. Sınıf - 2 Hafta",
    [
        ("Dijital Etik", ["Saygılı iletişim", "Telif hakkı", "Gizlilik"]),
        ("Güvenlik", ["Güçlü şifre", "Kişisel bilgi koruma", "Güvenli internet"]),
    ],
    "/home/user/BST/sunumlar/5-sinif-unite4-etik-guvenlik.pptx",
    (220, 50, 50)
)

# 5. Sınıf Ünite 5: Yapay Zeka
create_basic_ppt(
    "ÜNİTE 5: YAPAY ZEKA",
    "5. Sınıf - 4 Hafta",
    [
        ("YZ Nedir?", ["Yapay zeka tanımı", "Günlük hayatta YZ", "Örnekler"]),
        ("YZ Uygulamaları", ["Siri/Alexa", "Netflix önerileri", "Yüz tanıma"]),
    ],
    "/home/user/BST/sunumlar/5-sinif-unite5-yapay-zeka.pptx",
    (50, 150, 220)
)

# 5. Sınıf Ünite 6: Scratch Programlama
create_basic_ppt(
    "ÜNİTE 6: SCRATCH PROGRAMLAMA",
    "5. Sınıf - 14 Hafta",
    [
        ("Algoritma", ["Algoritma nedir?", "Günlük hayattan örnekler", "Akış diyagramı"]),
        ("Scratch Giriş", ["Scratch arayüzü", "Sprite", "Sahne", "Bloklar"]),
        ("Hareket", ["Hareket blokları", "Koordinatlar", "Dönme"]),
        ("Döngüler", ["Tekrar blokları", "Sonsuz döngü", "N kez tekrarla"]),
        ("Koşullar", ["If-Else", "Karar verme", "Karşılaştırma"]),
        ("Değişkenler", ["Değişken tanımlama", "Değer atama", "Kullanım"]),
        ("Projeler", ["Oyun yapımı", "Animasyon", "Hikaye anlatımı"]),
    ],
    "/home/user/BST/sunumlar/5-sinif-unite6-scratch.pptx",
    (220, 150, 50)
)

# 6. Sınıf Ünite 1: Bilişim Teknolojileri
create_basic_ppt(
    "ÜNİTE 1: BİLİŞİM TEKNOLOJİLERİ",
    "6. Sınıf - 4 Hafta",
    [
        ("Donanım-Yazılım", ["Temel kavramlar", "İşletim sistemleri", "Dosya yönetimi"]),
        ("İleri Konular", ["SSD vs HDD", "Bulut depolama", "Sanal makineler"]),
    ],
    "/home/user/BST/sunumlar/6-sinif-unite1-bilisim.pptx",
    (118, 75, 162)
)

# 6. Sınıf Ünite 2: Etik ve Güvenlik
create_basic_ppt(
    "ÜNİTE 2: ETİK VE GÜVENLİK",
    "6. Sınıf - 5 Hafta",
    [
        ("Dijital Etik", ["Telif hakkı", "Lisans türleri", "Dijital ayak izi"]),
        ("KVKK", ["Kişisel veri", "Haklar", "Sorumluluklar"]),
        ("Siber Tehditler", ["Phishing", "Virüs", "Fidye yazılımı"]),
        ("Güvenlik", ["2FA", "VPN", "HTTPS", "Güçlü şifre"]),
    ],
    "/home/user/BST/sunumlar/6-sinif-unite2-etik-guvenlik.pptx",
    (220, 50, 50)
)

# 6. Sınıf Ünite 3: İletişim ve İşbirliği
create_basic_ppt(
    "ÜNİTE 3: İLETİŞİM VE İŞBİRLİĞİ",
    "6. Sınıf - 2 Hafta",
    [
        ("Profesyonel E-posta", ["E-posta yapısı", "İş e-postası", "Netiquette"]),
        ("Bilgi Okuryazarlığı", ["Etkili arama", "CRAAP testi", "Atıf yapma"]),
    ],
    "/home/user/BST/sunumlar/6-sinif-unite3-iletisim.pptx",
    (100, 180, 220)
)

# 6. Sınıf Ünite 4: Ürün Oluşturma
create_basic_ppt(
    "ÜNİTE 4: ÜRÜN OLUŞTURMA",
    "6. Sınıf - 7 Hafta",
    [
        ("Grafik Tasarım", ["Tasarım ilkeleri", "Canva", "Renk teorisi"]),
        ("Video Düzenleme", ["Video editörleri", "Kırpma", "Geçişler", "Müzik"]),
        ("Animasyon", ["2D/3D animasyon", "Stop motion", "FPS"]),
        ("Podcast", ["Podcast nedir?", "Audacity", "Ses kaydı"]),
        ("Web Sitesi", ["Google Sites", "Web tasarım", "Yayınlama"]),
    ],
    "/home/user/BST/sunumlar/6-sinif-unite4-urun-olusturma.pptx",
    (180, 100, 220)
)

# 6. Sınıf Ünite 5: Problem Çözme
create_basic_ppt(
    "ÜNİTE 5: PROBLEM ÇÖZME VE PROGRAMLAMA",
    "6. Sınıf - 21 Hafta",
    [
        ("Problem Çözme", ["Anla-Planla-Uygula-Kontrol", "Algoritmik düşünme"]),
        ("İleri Scratch", ["Karmaşık projeler", "Fizik simülasyonu", "AI"]),
        ("Python Giriş", ["Python nedir?", "print() komutu", "Değişkenler"]),
        ("Projeler", ["Oyun geliştirme", "Simülasyonlar", "Final projesi"]),
    ],
    "/home/user/BST/sunumlar/6-sinif-unite5-problem-cozme.pptx",
    (220, 150, 50)
)

print("\n" + "="*50)
print("✅ TÜM ÜNİTE SUNUMLARI OLUŞTURULDU!")
print("="*50)
