#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
import os

def add_title_slide(prs, grade):
    """Kapak slaytı ekle"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # Başlık kutusu
    left = Inches(0.5)
    top = Inches(2)
    width = Inches(9)
    height = Inches(2)

    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True

    if grade == 5:
        p = tf.paragraphs[0]
        p.text = f"{grade}. SINIF"
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(54)
        p.font.bold = True
        p.font.color.rgb = RGBColor(102, 126, 234)

        p2 = tf.add_paragraph()
        p2.text = "Bilişim Teknolojileri ve Yazılım"
        p2.alignment = PP_ALIGN.CENTER
        p2.font.size = Pt(36)
        p2.font.color.rgb = RGBColor(102, 126, 234)

        p3 = tf.add_paragraph()
        p3.text = "FULL SUNUM - TÜM ÜNİTELER"
        p3.alignment = PP_ALIGN.CENTER
        p3.font.size = Pt(28)
        p3.font.color.rgb = RGBColor(118, 75, 162)
    else:
        p = tf.paragraphs[0]
        p.text = f"{grade}. SINIF"
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(54)
        p.font.bold = True
        p.font.color.rgb = RGBColor(79, 172, 254)

        p2 = tf.add_paragraph()
        p2.text = "Bilişim Teknolojileri ve Yazılım"
        p2.alignment = PP_ALIGN.CENTER
        p2.font.size = Pt(36)
        p2.font.color.rgb = RGBColor(79, 172, 254)

        p3 = tf.add_paragraph()
        p3.text = "FULL SUNUM - TÜM ÜNİTELER"
        p3.alignment = PP_ALIGN.CENTER
        p3.font.size = Pt(28)
        p3.font.color.rgb = RGBColor(0, 242, 254)

    print(f"✅ Kapak slaytı eklendi")


def copy_slide(source_prs, dest_prs, slide_index):
    """Bir slaytı başka bir sunuma kopyala"""
    source = source_prs.slides[slide_index]
    blank_layout = dest_prs.slide_layouts[6]  # Blank layout

    dest = dest_prs.slides.add_slide(blank_layout)

    # Tüm şekilleri kopyala
    for shape in source.shapes:
        el = shape.element
        newel = el.__class__(el)
        dest.shapes._spTree.insert_element_before(newel, 'p:extLst')

    # Notları kopyala
    if source.has_notes_slide:
        notes_slide = source.notes_slide
        if notes_slide.notes_text_frame.text:
            dest.notes_slide.notes_text_frame.text = notes_slide.notes_text_frame.text


def merge_presentations(input_files, output_file, grade):
    """Sunumları birleştir"""
    prs = Presentation()

    # Kapak ekle
    add_title_slide(prs, grade)

    total = 1

    for file_path in input_files:
        if not os.path.exists(file_path):
            print(f"⚠️  Bulunamadı: {os.path.basename(file_path)}")
            continue

        print(f"📄 Ekleniyor: {os.path.basename(file_path)}")

        try:
            source_prs = Presentation(file_path)

            for i in range(len(source_prs.slides)):
                try:
                    copy_slide(source_prs, prs, i)
                    total += 1
                except Exception as e:
                    print(f"  ⚠️  Slayt {i+1} kopyalanamadı: {e}")

            print(f"  ✅ {len(source_prs.slides)} slayt eklendi")

        except Exception as e:
            print(f"  ❌ Hata: {e}")

    # Kaydet
    prs.save(output_file)
    print(f"\n🎉 {total} slayt kaydedildi: {os.path.basename(output_file)}\n")

    return total


# 5. Sınıf
print("=" * 60)
print("5. SINIF FULL SUNUM")
print("=" * 60)

grade5_files = [
    '/home/user/BST/sunumlar/5-sinif-unite1-bilisim-temelleri-DETAYLI.pptx',
    '/home/user/BST/sunumlar/5-sinif-unite2-dijital-urun-tasarimi-DETAYLI.pptx',
    '/home/user/BST/sunumlar/5-sinif-unite3-aglar-iletisim-DETAYLI.pptx',
    '/home/user/BST/sunumlar/5-sinif-unite4-etik-guvenlik-DETAYLI.pptx',
    '/home/user/BST/sunumlar/5-sinif-unite5-yapay-zeka-DETAYLI.pptx',
    '/home/user/BST/sunumlar/5-sinif-unite6-scratch-DETAYLI.pptx'
]

total_5 = merge_presentations(
    grade5_files,
    '/home/user/BST/sunumlar/5-SINIF-FULL-SUNUM-DETAYLI.pptx',
    5
)

# 6. Sınıf
print("=" * 60)
print("6. SINIF FULL SUNUM")
print("=" * 60)

grade6_files = [
    '/home/user/BST/sunumlar/6-sinif-unite1-veri-analiz-DETAYLI.pptx',
    '/home/user/BST/sunumlar/6-sinif-unite2-kelime-islemci-DETAYLI.pptx',
    '/home/user/BST/sunumlar/6-sinif-unite3-algoritmalar-DETAYLI.pptx',
    '/home/user/BST/sunumlar/6-sinif-unite4-scratch-ileri-DETAYLI.pptx',
    '/home/user/BST/sunumlar/6-sinif-unite5-sunum-programlari-DETAYLI.pptx'
]

total_6 = merge_presentations(
    grade6_files,
    '/home/user/BST/sunumlar/6-SINIF-FULL-SUNUM-DETAYLI.pptx',
    6
)

print("=" * 60)
print("TAMAMLANDI")
print("=" * 60)
print(f"✅ 5. Sınıf: {total_5} slayt")
print(f"✅ 6. Sınıf: {total_6} slayt")
