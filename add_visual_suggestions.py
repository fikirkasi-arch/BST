#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from pptx import Presentation
import os
import re

def get_slide_text(slide):
    """Slayttaki tüm metni al"""
    text_content = []
    for shape in slide.shapes:
        if hasattr(shape, "text"):
            text_content.append(shape.text.lower())
    return " ".join(text_content)

def suggest_visuals(slide_text, unit_topic):
    """Slayt içeriğine göre görsel/video önerileri oluştur"""
    suggestions = []

    # Donanım konuları
    if any(word in slide_text for word in ['cpu', 'işlemci', 'ram', 'bellek']):
        suggestions.append("📸 Görsel: CPU, RAM ve anakart fotoğrafları")
        suggestions.append("🎬 Video: Bilgisayar donanımı tanıtım videosu (1-2 dk)")
        suggestions.append("💡 İnfografik: Bilgisayar bileşenleri şeması")

    if any(word in slide_text for word in ['klavye', 'fare', 'mouse', 'giriş aygıt']):
        suggestions.append("📸 Görsel: Çeşitli giriş aygıtları (klavye, fare, mikrofon, kamera)")
        suggestions.append("🎬 Video: Giriş aygıtlarının nasıl çalıştığını gösteren animasyon")

    if any(word in slide_text for word in ['monitör', 'yazıcı', 'hoparlör', 'çıkış aygıt']):
        suggestions.append("📸 Görsel: Çeşitli çıkış aygıtları fotoğrafları")
        suggestions.append("🎬 Video: Monitör ve yazıcı çalışma prensibi")

    if any(word in slide_text for word in ['hard disk', 'ssd', 'usb', 'depolama']):
        suggestions.append("📸 Görsel: HDD vs SSD karşılaştırması")
        suggestions.append("🎬 Video: Depolama birimleri animasyonu")
        suggestions.append("💡 İnfografik: Depolama kapasiteleri karşılaştırma tablosu")

    # Yazılım konuları
    if any(word in slide_text for word in ['işletim sistemi', 'windows', 'linux', 'macos']):
        suggestions.append("📸 Görsel: Farklı işletim sistemleri logoları ve ekran görüntüleri")
        suggestions.append("🎬 Video: İşletim sistemi nedir? Animasyonlu anlatım")

    if any(word in slide_text for word in ['word', 'kelime işlemci', 'metin düzenle']):
        suggestions.append("📸 Görsel: MS Word arayüz ekran görüntüsü")
        suggestions.append("🎬 Video: Word temel özellikleri demo (2-3 dk)")
        suggestions.append("💡 Ekran kaydı: Biçimlendirme araçları kullanımı")

    # Ağ ve internet
    if any(word in slide_text for word in ['ağ', 'network', 'lan', 'wan', 'internet']):
        suggestions.append("📸 Görsel: LAN ve WAN ağ topolojileri şeması")
        suggestions.append("🎬 Video: İnternet nasıl çalışır? Animasyonlu anlatım")
        suggestions.append("💡 İnfografik: Ev ağı kurulumu diyagramı")

    if any(word in slide_text for word in ['router', 'modem', 'switch', 'access point']):
        suggestions.append("📸 Görsel: Router, modem ve switch cihazları")
        suggestions.append("🎬 Video: Modem ve router farkı")

    if any(word in slide_text for word in ['ip adresi', 'dns', 'protokol']):
        suggestions.append("📸 Görsel: IP adresi formatı ve yapısı")
        suggestions.append("🎬 Video: DNS nasıl çalışır? Animasyon")
        suggestions.append("💡 İnfografik: TCP/IP protokol katmanları")

    # Güvenlik konuları
    if any(word in slide_text for word in ['güvenlik', 'şifre', 'password', 'parola']):
        suggestions.append("📸 Görsel: Güçlü şifre örnekleri vs zayıf şifre")
        suggestions.append("🎬 Video: Şifre güvenliği önemi (1-2 dk)")
        suggestions.append("💡 İnfografik: Güçlü şifre oluşturma kuralları")

    if any(word in slide_text for word in ['virüs', 'trojan', 'malware', 'zararlı']):
        suggestions.append("📸 Görsel: Farklı zararlı yazılım türleri ikonları")
        suggestions.append("🎬 Video: Virüslerden korunma yolları")

    if any(word in slide_text for word in ['phishing', 'kimlik avı', 'dolandırıcılık']):
        suggestions.append("📸 Görsel: Phishing e-posta örnekleri (gerçek vakalar)")
        suggestions.append("🎬 Video: Kimlik avı saldırılarından korunma")
        suggestions.append("💡 Ekran görüntüsü: Sahte web sitesi tespiti")

    if any(word in slide_text for word in ['firewall', 'güvenlik duvarı', 'antivir']):
        suggestions.append("📸 Görsel: Firewall çalışma prensibi şeması")
        suggestions.append("🎬 Video: Antivirüs programı nasıl çalışır?")

    # Etik ve telif
    if any(word in slide_text for word in ['telif', 'copyright', 'fikri mülkiyet']):
        suggestions.append("📸 Görsel: Telif hakları sembolleri (©, ®, ™)")
        suggestions.append("🎬 Video: Dijital etik ve telif hakları")
        suggestions.append("💡 İnfografik: Lisans türleri karşılaştırması")

    if any(word in slide_text for word in ['kvkk', 'kişisel veri', 'gizlilik']):
        suggestions.append("📸 Görsel: KVKK logosu ve temel ilkeler")
        suggestions.append("🎬 Video: Kişisel veri nedir? Animasyon")

    # Yapay Zeka
    if any(word in slide_text for word in ['yapay zeka', 'ai', 'artificial intelligence']):
        suggestions.append("📸 Görsel: Yapay zeka uygulama örnekleri (Siri, Alexa, ChatGPT)")
        suggestions.append("🎬 Video: Yapay zeka nedir? Çocuklar için anlatım")
        suggestions.append("💡 İnfografik: AI kullanım alanları")

    if any(word in slide_text for word in ['makine öğrenmesi', 'machine learning']):
        suggestions.append("📸 Görsel: Makine öğrenmesi süreci diyagramı")
        suggestions.append("🎬 Video: Makine öğrenmesi basit örneklerle")

    if any(word in slide_text for word in ['chatbot', 'sohbet robot', 'ses tanıma']):
        suggestions.append("📸 Görsel: Popüler chatbot örnekleri (Siri, Google Assistant)")
        suggestions.append("🎬 Video: Chatbot nasıl çalışır?")

    # Scratch ve Kodlama
    if any(word in slide_text for word in ['scratch', 'blok', 'kod', 'programla']):
        suggestions.append("📸 Görsel: Scratch arayüzü ekran görüntüsü")
        suggestions.append("🎬 Video: Scratch ile basit oyun yapımı (3-5 dk)")
        suggestions.append("💡 Ekran kaydı: Adım adım Scratch projesi")

    if any(word in slide_text for word in ['algoritma', 'akış', 'adım']):
        suggestions.append("📸 Görsel: Akış şeması örnekleri")
        suggestions.append("🎬 Video: Algoritma nedir? Günlük hayattan örnekler")
        suggestions.append("💡 İnfografik: Akış şeması sembolleri")

    if any(word in slide_text for word in ['döngü', 'loop', 'tekrar']):
        suggestions.append("📸 Görsel: Döngü mantığı şeması")
        suggestions.append("🎬 Video: Döngüler Scratch'te nasıl kullanılır?")

    if any(word in slide_text for word in ['koşul', 'if', 'karar']):
        suggestions.append("📸 Görsel: Koşullu yapılar akış diyagramı")
        suggestions.append("🎬 Video: Koşullu ifadeler Scratch'te")

    if any(word in slide_text for word in ['değişken', 'variable', 'veri saklama']):
        suggestions.append("📸 Görsel: Değişken kavramı görselleştirmesi")
        suggestions.append("🎬 Video: Scratch'te değişken kullanımı")

    # Veri ve analiz
    if any(word in slide_text for word in ['veri', 'data', 'tablo', 'grafik']):
        suggestions.append("📸 Görsel: Farklı grafik türleri örnekleri")
        suggestions.append("🎬 Video: Veri görselleştirme önemi")
        suggestions.append("💡 İnfografik: Grafik türleri ve kullanım alanları")

    if any(word in slide_text for word in ['excel', 'hesap tablosu', 'hücre']):
        suggestions.append("📸 Görsel: Excel arayüzü ekran görüntüsü")
        suggestions.append("🎬 Video: Excel temel işlemler (2-3 dk)")
        suggestions.append("💡 Ekran kaydı: Formül kullanımı demo")

    # Dijital ürün tasarımı
    if any(word in slide_text for word in ['paint', 'çizim', 'resim', 'boyama']):
        suggestions.append("📸 Görsel: Paint araçları ekran görüntüsü")
        suggestions.append("🎬 Video: Paint ile basit çizim teknikleri")

    if any(word in slide_text for word in ['tasarım', 'design', 'renk', 'şekil']):
        suggestions.append("📸 Görsel: Temel tasarım prensipleri")
        suggestions.append("🎬 Video: Dijital tasarıma giriş")
        suggestions.append("💡 İnfografik: Renk teorisi ve uyumu")

    # Sunum programları
    if any(word in slide_text for word in ['powerpoint', 'sunum', 'slayt', 'presentation']):
        suggestions.append("📸 Görsel: PowerPoint arayüzü ve araçları")
        suggestions.append("🎬 Video: Etkili sunum hazırlama ipuçları")
        suggestions.append("💡 Ekran kaydı: PowerPoint temel özellikler")

    # Genel teknoloji
    if any(word in slide_text for word in ['bilişim', 'teknoloji', 'dijital']):
        suggestions.append("📸 Görsel: Dijital çağ görselleri ve teknoloji örnekleri")
        suggestions.append("🎬 Video: Bilişim teknolojilerinin tarihçesi")

    # Eğer hiç öneri yoksa genel öneriler ekle
    if not suggestions:
        suggestions.append("📸 Görsel: Konuyla ilgili görsel içerik eklenebilir")
        suggestions.append("💡 İlgili ekran görüntüleri veya şemalar eklenebilir")

    return suggestions

def add_notes_to_presentations():
    """Tüm DETAYLI sunumlara notlar ekle"""

    presentations = [
        ('5-sinif-unite1-bilisim-temelleri-DETAYLI.pptx', 'Bilişim Temelleri'),
        ('5-sinif-unite2-dijital-urun-tasarimi-DETAYLI.pptx', 'Dijital Ürün Tasarımı'),
        ('5-sinif-unite3-aglar-iletisim-DETAYLI.pptx', 'Ağlar ve İletişim'),
        ('5-sinif-unite4-etik-guvenlik-DETAYLI.pptx', 'Etik ve Güvenlik'),
        ('5-sinif-unite5-yapay-zeka-DETAYLI.pptx', 'Yapay Zeka'),
        ('5-sinif-unite6-scratch-DETAYLI.pptx', 'Scratch Programlama'),
        ('6-sinif-unite1-veri-analiz-DETAYLI.pptx', 'Veri Analizi'),
        ('6-sinif-unite2-kelime-islemci-DETAYLI.pptx', 'Kelime İşlemci'),
        ('6-sinif-unite3-algoritmalar-DETAYLI.pptx', 'Algoritmalar'),
        ('6-sinif-unite4-scratch-ileri-DETAYLI.pptx', 'Scratch İleri Seviye'),
        ('6-sinif-unite5-sunum-programlari-DETAYLI.pptx', 'Sunum Programları'),
    ]

    total_slides = 0
    total_notes_added = 0

    for filename, unit_topic in presentations:
        filepath = f'/home/user/BST/sunumlar/{filename}'

        if not os.path.exists(filepath):
            print(f"⚠️  Dosya bulunamadı: {filename}")
            continue

        print(f"\n{'='*60}")
        print(f"📄 İşleniyor: {filename}")
        print(f"{'='*60}")

        try:
            prs = Presentation(filepath)
            slides_with_notes = 0

            for slide_num, slide in enumerate(prs.slides, 1):
                # Slayt metnini al
                slide_text = get_slide_text(slide)

                # Görsel önerileri oluştur
                suggestions = suggest_visuals(slide_text, unit_topic)

                if suggestions:
                    # Mevcut notları koru
                    existing_notes = ""
                    if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text:
                        existing_notes = slide.notes_slide.notes_text_frame.text + "\n\n"

                    # Yeni notları ekle
                    notes_text = existing_notes + "🎨 GÖRSELLEŞTİRME ÖNERİLERİ:\n\n"
                    notes_text += "\n".join(suggestions)

                    # Notları slayta ekle
                    slide.notes_slide.notes_text_frame.text = notes_text

                    slides_with_notes += 1
                    total_notes_added += len(suggestions)

                    print(f"  ✅ Slayt {slide_num}: {len(suggestions)} öneri eklendi")

            # Sunumu kaydet
            prs.save(filepath)
            total_slides += len(prs.slides)

            print(f"\n  💾 Kaydedildi: {slides_with_notes}/{len(prs.slides)} slayta not eklendi")

        except Exception as e:
            print(f"  ❌ Hata: {e}")
            continue

    print(f"\n{'='*60}")
    print(f"ÖZET")
    print(f"{'='*60}")
    print(f"✅ Toplam {total_slides} slayt işlendi")
    print(f"📝 Toplam {total_notes_added} görsel/video önerisi eklendi")
    print(f"💾 Tüm sunumlar güncellendi")

if __name__ == '__main__':
    add_notes_to_presentations()
