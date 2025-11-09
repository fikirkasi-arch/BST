#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from pptx import Presentation
import os

def verify_presentation(prs_path):
    """Sunumun içeriğini doğrula"""

    prs = Presentation(prs_path)
    filename = os.path.basename(prs_path)

    print(f"\n{'='*70}")
    print(f"📄 {filename}")
    print(f"{'='*70}")

    total_slides = len(prs.slides)
    empty_slides = 0
    slides_with_content = 0

    for slide_num, slide in enumerate(prs.slides, 1):
        has_text = False
        has_shapes = False

        # Metin var mı?
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                has_text = True
            has_shapes = True

        if has_text:
            slides_with_content += 1
        elif has_shapes:
            slides_with_content += 1  # Şekil var ama metin yok (diagram vb)
        else:
            empty_slides += 1
            print(f"  ⚠️  Slayt {slide_num}: BOŞ (metin veya şekil yok)")

    print(f"\n  📊 Toplam slayt: {total_slides}")
    print(f"  ✅ İçerikli slayt: {slides_with_content}")
    print(f"  ⚠️  Boş slayt: {empty_slides}")

    if empty_slides == 0:
        print(f"  🎉 TÜM SLAYTLAR DOLU!")

    return empty_slides == 0

def main():
    """FULL sunumları doğrula"""

    presentations = [
        '/home/user/BST/sunumlar/5-SINIF-FULL-SUNUM-DETAYLI.pptx',
        '/home/user/BST/sunumlar/6-SINIF-FULL-SUNUM-DETAYLI.pptx',
    ]

    print("="*70)
    print("FULL SUNUMLARI DOĞRULAMA")
    print("="*70)

    all_valid = True

    for prs_path in presentations:
        if not os.path.exists(prs_path):
            print(f"⚠️  Dosya bulunamadı: {prs_path}")
            continue

        try:
            is_valid = verify_presentation(prs_path)
            if not is_valid:
                all_valid = False
        except Exception as e:
            print(f"❌ Hata: {e}")
            all_valid = False

    print(f"\n{'='*70}")
    if all_valid:
        print("✅ TÜM SUNUMLAR GEÇERLİ - SLAYTLAR DOLU")
    else:
        print("⚠️  BAZI SUNUMLARDA BOŞ SLAYTLAR VAR")
    print("="*70)

if __name__ == '__main__':
    main()
