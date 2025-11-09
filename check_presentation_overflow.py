#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from pptx import Presentation
from pptx.util import Pt
import os

def check_text_overflow(prs_path):
    """Sunum slaytlarında metin taşmalarını kontrol et"""

    prs = Presentation(prs_path)
    filename = os.path.basename(prs_path)

    print(f"\n{'='*70}")
    print(f"📄 {filename}")
    print(f"{'='*70}")

    issues_found = False

    for slide_num, slide in enumerate(prs.slides, 1):
        slide_issues = []

        for shape_num, shape in enumerate(slide.shapes):
            if not hasattr(shape, "text_frame"):
                continue

            text_frame = shape.text_frame

            # Toplam karakter sayısını kontrol et
            total_chars = sum(len(p.text) for p in text_frame.paragraphs)

            # Çok uzun metin uyarısı
            if total_chars > 500:
                slide_issues.append(f"  ⚠️  Şekil {shape_num+1}: ÇOK UZUN METİN ({total_chars} karakter)")
                issues_found = True
            elif total_chars > 300:
                slide_issues.append(f"  ⚡ Şekil {shape_num+1}: Uzun metin ({total_chars} karakter)")
                issues_found = True

            # Paragraf başına kontrol
            for para_num, paragraph in enumerate(text_frame.paragraphs):
                # Çok uzun paragraf
                if len(paragraph.text) > 200:
                    slide_issues.append(f"  ⚠️  Şekil {shape_num+1}, Paragraf {para_num+1}: Çok uzun ({len(paragraph.text)} karakter)")
                    issues_found = True

                # Küçük font kontrolü
                for run in paragraph.runs:
                    if run.font.size and run.font.size < Pt(10):
                        slide_issues.append(f"  👓 Şekil {shape_num+1}: Küçük font ({run.font.size.pt}pt)")
                        issues_found = True
                        break

        # Slayt için sorun varsa yazdır
        if slide_issues:
            print(f"\nSlayt {slide_num}:")
            for issue in slide_issues:
                print(issue)

    if not issues_found:
        print(f"\n✅ Sorun tespit edilmedi - Tüm içerikler uygun boyutta")

    return issues_found

def main():
    """Tüm DETAYLI sunumları kontrol et"""

    presentations = [
        '5-sinif-unite1-bilisim-temelleri-DETAYLI.pptx',
        '5-sinif-unite2-dijital-urun-tasarimi-DETAYLI.pptx',
        '5-sinif-unite3-aglar-iletisim-DETAYLI.pptx',
        '5-sinif-unite4-etik-guvenlik-DETAYLI.pptx',
        '5-sinif-unite5-yapay-zeka-DETAYLI.pptx',
        '5-sinif-unite6-scratch-DETAYLI.pptx',
        '6-sinif-unite1-veri-analiz-DETAYLI.pptx',
        '6-sinif-unite2-kelime-islemci-DETAYLI.pptx',
        '6-sinif-unite3-algoritmalar-DETAYLI.pptx',
        '6-sinif-unite4-scratch-ileri-DETAYLI.pptx',
        '6-sinif-unite5-sunum-programlari-DETAYLI.pptx',
    ]

    total_issues = 0

    for filename in presentations:
        filepath = f'/home/user/BST/sunumlar/{filename}'

        if not os.path.exists(filepath):
            print(f"⚠️  Dosya bulunamadı: {filename}")
            continue

        try:
            if check_text_overflow(filepath):
                total_issues += 1
        except Exception as e:
            print(f"❌ Hata: {e}")

    print(f"\n{'='*70}")
    print(f"ÖZET")
    print(f"{'='*70}")

    if total_issues == 0:
        print(f"✅ Tüm sunumlar kontrol edildi - Ciddi sorun tespit edilmedi")
    else:
        print(f"⚠️  {total_issues} sunumda potansiyel taşma sorunları tespit edildi")
        print(f"💡 Not: Tespit edilen sorunların çoğu otomatik olarak PowerPoint")
        print(f"   tarafından halledilebilir. Manuel kontrol önerilir.")

    print(f"\n📝 ÖNERİLER:")
    print(f"  • Çok uzun metinler için bullet point kullanın")
    print(f"  • Font boyutunu 12-18pt arasında tutun")
    print(f"  • Slayt başına maksimum 6-7 madde kullanın")
    print(f"  • Detaylı açıklamaları notlar bölümüne ekleyin")

if __name__ == '__main__':
    main()
